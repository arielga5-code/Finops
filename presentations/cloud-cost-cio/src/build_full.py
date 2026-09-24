from pptx import Presentation
from pptx.chart.data import CategoryChartData

SRC = "deck/Harel_Cloud_Cost_CIO_Final.pptx"
DST = "deck/Harel_Cloud_Cost_CIO_August.pptx"
#        main 9                        appendix 8
KEEP = [1, 10, 3, 2, 5, 9, 12, 32, 28, 29, 30, 31, 11, 33, 34, 35, 36]
M8 = ['Jan','Feb','Mar','Apr','May','Jun','Jul','Aug']

def surgery(src, dst, keep):
    p = Presentation(src); lst = p.slides._sldIdLst; ids = list(lst)
    k0 = [k-1 for k in keep]
    for i, sid in enumerate(ids):
        if i not in k0:
            p.part.drop_rel(sid.rId); lst.remove(sid)
    rem = {i: sid for i, sid in enumerate(ids) if i in k0}
    for sid in list(lst): lst.remove(sid)
    for i in k0: lst.append(rem[i])
    p.save(dst)

def T(slide, idx, text):
    tf = slide.shapes[idx].text_frame
    para = tf.paragraphs[0]
    if not para.runs: para.add_run()
    para.runs[0].text = text
    for r in para.runs[1:]: r._r.getparent().remove(r._r)
    for e in tf.paragraphs[1:]: e._p.getparent().remove(e._p)

def cell(t, r, c, text):
    para = t.cell(r, c).text_frame.paragraphs[0]
    if not para.runs: para.add_run()
    para.runs[0].text = text
    for x in para.runs[1:]: x._r.getparent().remove(x._r)

def charts(slide):
    return [sh.chart for sh in slide.shapes if getattr(sh, 'has_chart', False)]
def table_of(slide):
    for sh in slide.shapes:
        if getattr(sh, 'has_table', False): return sh.table

def setchart(ch, cats, series):
    cd = CategoryChartData(); cd.categories = cats
    for n, v in series: cd.add_series(n, v)
    ch.replace_data(cd)

surgery(SRC, DST, KEEP)
p = Presentation(DST); S = p.slides
exec(open('slides_main.py').read())
exec(open('slides_appendix.py').read())
p.save(DST)
print("built", DST, "with", len(S), "slides")
